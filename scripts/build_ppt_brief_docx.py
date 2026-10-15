"""将已核实的15页答辩提纲制成供PPT生成工具读取的Word制作稿。"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "materials" / "答辩PPT内容.md"
OUTPUT = ROOT / "materials" / "基于YOLOv8的远景行人检测_PPT制作稿.docx"
FIGURES = ROOT / "materials" / "figures"

NAVY = RGBColor(25, 52, 88)
TEAL = RGBColor(0, 112, 126)
GRAY = RGBColor(75, 85, 96)


def set_font(style, size, color=None, bold=False):
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia"):
        rfonts.set(qn(f"w:{key}"), "Microsoft YaHei")


def clean(text):
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("  ", " ")
    text = text.replace("$$", "").replace("\\longrightarrow", "→")
    text = text.replace("\\times", "×").replace("\\geq", "≥")
    return text.strip()


def shade(cell, fill="EEF4F6"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def add_table(doc, rows):
    data = [[clean(c) for c in row.strip().strip("|").split("|")] for row in rows]
    data = [r for r in data if not all(re.fullmatch(r"[-: ]+", c or "") for c in r)]
    if not data:
        return
    table = doc.add_table(rows=len(data), cols=max(map(len, data)))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, row in enumerate(data):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                shade(cell, "DDEBEF")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True
                        run.font.color.rgb = NAVY
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_labeled(doc, label, value, color=GRAY):
    if not value:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    lead = p.add_run(label + "：")
    lead.bold = True
    lead.font.color.rgb = color
    p.add_run(clean(value))


def add_image(doc, name, caption, width=5.4):
    path = FIGURES / name
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph(clean(caption))
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.style = "Caption"


def add_pair(doc, left, right, caption):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, name in zip(table.rows[0].cells, (left, right)):
        cell.width = Inches(3.1)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(FIGURES / name), width=Inches(3.05))
    cp = doc.add_paragraph(caption)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.style = "Caption"


def parse_slides(text):
    chunks = re.split(r"(?=^## 第\d+页 )", text, flags=re.M)
    slides = []
    for chunk in chunks:
        m = re.match(r"^## 第(\d+)页 (.+)$", chunk, flags=re.M)
        if not m:
            continue
        fields = {"页面内容": [], "视觉": [], "讲解备注": [], "来源": [], "明显标注": []}
        current = None
        for line in chunk.splitlines()[1:]:
            marker = re.match(r"^\*\*(页面内容|视觉|讲解备注|来源|明显标注)\*\*[：:]?\s*(.*)$", line)
            if marker:
                current = marker.group(1)
                if marker.group(2).strip():
                    fields[current].append(marker.group(2).strip())
            elif current is not None:
                fields[current].append(line)
        slides.append((int(m.group(1)), m.group(2).strip(), fields))
    return slides


def add_content(doc, lines):
    table_lines = []
    for line in [*lines, ""]:
        if line.strip().startswith("|"):
            table_lines.append(line)
            continue
        if table_lines:
            add_table(doc, table_lines)
            table_lines = []
        line = clean(line)
        if not line:
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(line[2:], style="List Bullet")
        else:
            p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(4)


def main():
    slides = parse_slides(SOURCE.read_text(encoding="utf-8"))
    if [n for n, _, _ in slides] != list(range(1, 16)):
        raise ValueError("答辩源稿必须恰好包含第1至15页。")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.5)
    sec.left_margin, sec.right_margin = Cm(1.8), Cm(1.8)
    set_font(doc.styles["Normal"], 10, GRAY)
    set_font(doc.styles["Title"], 22, NAVY, True)
    set_font(doc.styles["Heading 1"], 16, NAVY, True)
    set_font(doc.styles["Heading 2"], 11, TEAL, True)
    set_font(doc.styles["Caption"], 8.5, GRAY)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.2
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("基于YOLOv8的远景行人检测")
    p = doc.add_paragraph("15页答辩PPT制作稿｜供豆包生成演示文稿", style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("本稿按一页Word对应一页幻灯片组织，配图已内嵌。姓名、学号、班级及实际分工须由学生填写；演示录屏尚未提供。")

    doc.add_heading("请按以下要求制作", level=1)
    for item in (
        "制作恰好15页中文16:9演示文稿，封面算第1页；白底、深蓝标题、黑色正文，少字大图。",
        "保留每页的讲解备注；图表和模型结构图尽量用可编辑对象制作，已内嵌的实测曲线和检测图片保持原样。",
        "所有实验数值必须逐项照录，不得把验证集数值写成test-dev结果，也不得声称ECA+P2优于单独P2。",
        "第14页预留1–2分钟真实运行录屏的位置；目前只能使用静态图片并标注“图片检测展示”。",
        "P/R采用置信度0.25、IoU0.5；AP_small按原图人体框面积小于32²像素；速度为RTX PRO 5000上FP32、batch=1、50次预热及500次同步计时。",
    ):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("必须保留的结果", level=1)
    add_table(doc, [
        "| 模型 | test-dev mAP50–95 | AP_small | 端到端延迟ms |",
        "|---|---:|---:|---:|",
        "| baseline | 14.78% | 11.08% | 6.78 |",
        "| eca | 15.01% | 11.04% | 6.61 |",
        "| p2 | 17.35% | 13.43% | 7.25 |",
        "| eca_p2 | 17.28% | 13.45% | 7.91 |",
    ])
    doc.add_paragraph("结果来源：项目 results_bundle_e100_complete.zip；四组各训练100轮，960输入、batch=8、seed=42。")

    for number, title, fields in slides:
        doc.add_page_break()
        doc.add_heading(f"第{number}页  {title}", level=1)
        doc.add_heading("页面展示文字", level=2)
        add_content(doc, fields["页面内容"])
        if fields["明显标注"]:
            add_labeled(doc, "页面必须标注", " ".join(fields["明显标注"]), TEAL)

        if number == 3:
            add_image(doc, "raw_9999938_00000_d_0000207.jpg", "图：VisDrone2019-DET test-dev原始航拍图像；这是原图，不是检测输出。", 5.4)
        elif number == 11:
            add_image(doc, "eca_p2_threshold_curve.png", "图：完整模型验证集阈值曲线。", 4.4)
        elif number == 12:
            add_image(doc, "eca_p2_training.png", "图：完整模型100轮训练与验证曲线；本页流程检查数字与正式实验分开。", 5.3)
        elif number == 13:
            add_image(doc, "eca_p2_pr_curve.png", "图：完整模型验证集IoU0.5 PR曲线；上表数字来自独立test-dev。", 4.1)
        elif number == 14:
            add_pair(doc, "baseline_9999938_00000_d_0000207.jpg", "eca_p2_9999938_00000_d_0000207.jpg", "同一test-dev图像：左为基线，右为ECA+P2；绿框GT、红框预测，阈值0.25。")

        add_labeled(doc, "配图与版式", " ".join(fields["视觉"]))
        add_labeled(doc, "讲解备注", " ".join(fields["讲解备注"]))
        add_labeled(doc, "来源", " ".join(fields["来源"]))

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("基于YOLOv8的远景行人检测｜答辩PPT制作稿")
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = GRAY
    doc.core_properties.title = "基于YOLOv8的远景行人检测 PPT制作稿"
    doc.core_properties.subject = "15页中文答辩PPT内容与配图"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
